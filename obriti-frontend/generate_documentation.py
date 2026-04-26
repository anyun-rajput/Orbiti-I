#!/usr/bin/env python3
"""
Generate comprehensive DOCX documentation for Orbit-I Frontend
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

def add_heading(doc, text, level=1, color=None):
    """Add a heading to the document"""
    heading = doc.add_heading(text, level=level)
    if color:
        for run in heading.runs:
            run.font.color.rgb = color
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return heading

def add_paragraph(doc, text, bold=False, italic=False, size=11):
    """Add a paragraph to the document"""
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
    return p

def add_code_block(doc, code, language="typescript"):
    """Add a code block with monospace font"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)
    p.style = 'Intense Quote'
    return p

def create_table(doc, rows, cols, header_data=None):
    """Create a table in the document"""
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Light Grid Accent 1'
    if header_data:
        for i, header in enumerate(header_data):
            table.rows[0].cells[i].text = header
    return table

# Create document
doc = Document()

# Add title
title = doc.add_heading('ORBIT-I FRONTEND', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.runs[0]
title_run.font.color.rgb = RGBColor(0, 51, 102)
title_run.font.size = Pt(28)

# Add subtitle
subtitle = doc.add_heading('Comprehensive Technical Documentation', level=2)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add metadata
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_run = meta.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y')}\nVersion: 1.0")
meta_run.font.size = Pt(10)
meta_run.font.italic = True

doc.add_paragraph()  # Spacing

# Table of Contents
add_heading(doc, "TABLE OF CONTENTS", level=1)
toc_items = [
    "1. Executive Summary",
    "2. Project Overview",
    "3. Architecture & Design Patterns",
    "4. Technology Stack",
    "5. Project Structure",
    "6. Authentication & Authorization",
    "7. Core Components",
    "8. Pages & Routes",
    "9. Services & API Integration",
    "10. State Management",
    "11. UI Components Library",
    "12. Styling & Theming",
    "13. Utilities & Hooks",
    "14. Development Workflow",
    "15. Build & Deployment",
    "16. Security Considerations",
    "17. Performance Optimization",
    "18. Testing Strategy",
    "19. Troubleshooting & FAQ",
]
for item in toc_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 1. Executive Summary
add_heading(doc, "1. EXECUTIVE SUMMARY", level=1)
add_paragraph(doc, """
Orbit-I is an enterprise vulnerability management platform frontend built with modern web technologies. 
It provides a comprehensive dashboard for security teams to manage vulnerabilities, assets, exceptions, 
and integrations. The application emphasizes security, role-based access control, and user experience.

Key Characteristics:
• Modern React with TypeScript for type safety
• Vite for fast development and optimized builds
• Role-based and permission-based access control
• Real-time data with React Query (TanStack Query)
• Responsive UI with Tailwind CSS and Shadcn/UI components
• Docker containerization for deployment
• Microsoft SSO integration support
• Comment system for vulnerability collaboration
""")

doc.add_page_break()

# 2. Project Overview
add_heading(doc, "2. PROJECT OVERVIEW", level=1)

add_heading(doc, "2.1 Purpose & Goals", level=2)
add_paragraph(doc, """
Orbit-I is designed to help organizations:
• Monitor and manage security vulnerabilities across their infrastructure
• Track assets and their security status
• Manage exceptions and remediation workflows
• Collaborate on vulnerability resolution through comments and discussions
• Control access through role-based and permission-based authorization
• Integrate with third-party platforms (Microsoft SSO, etc.)
""")

add_heading(doc, "2.2 Target Users", level=2)
add_paragraph(doc, """
• Security Teams: Monitor vulnerabilities and manage remediation
• System Administrators: Manage assets and rescan workflows
• Organization Managers: View reports and manage exceptions
• Integrations Admins: Configure third-party service integrations
""")

add_heading(doc, "2.3 Key Features", level=2)
add_paragraph(doc, """
• Dashboard: Real-time vulnerability overview with charts and analytics
• Vulnerability Management: Search, filter, and analyze vulnerabilities
• Asset Management: Track and manage infrastructure assets
• Exception Management: Handle approved vulnerability exceptions
• User Management: Control user roles and permissions
• Integrations: Connect with external services (Microsoft AAD, etc.)
• Rescan Management: Organize and track vulnerability rescans
• Comments & Collaboration: Add comments to vulnerabilities and hosts
• Role-Based Access Control (RBAC): Admin, team, and limited user roles
""")

doc.add_page_break()

# 3. Architecture & Design Patterns
add_heading(doc, "3. ARCHITECTURE & DESIGN PATTERNS", level=1)

add_heading(doc, "3.1 Architectural Overview", level=2)
add_paragraph(doc, """
Orbit-I follows a modern Single Page Application (SPA) architecture with clear separation of concerns:

Frontend (This Application)
├── Presentation Layer: React components and pages
├── Business Logic Layer: Services and hooks
└── Data Layer: API integration and state management

Backend API (Separate Service)
└── RESTful API endpoints for data persistence
""")

add_heading(doc, "3.2 Design Patterns Used", level=2)

add_heading(doc, "Component Pattern", level=3)
add_paragraph(doc, """
• Functional Components: All components use React hooks
• Container/Presentational: Smart components fetch data, dumb components display
• Higher-Order Components (HOC): Used for route protection (ProtectedRoute, AdminRoute, PermissionRoute)
""")

add_heading(doc, "State Management Pattern", level=3)
add_paragraph(doc, """
• React Context API: Global auth state (AuthProvider)
• React Query: Server state management and caching
• Local State: Component-level state with useState
""")

add_heading(doc, "Service Pattern", level=3)
add_paragraph(doc, """
• Service Classes: Encapsulate API logic (MicrosoftSSOService, SecureTokenManager)
• Custom Hooks: Reusable logic (use-toast, use-mobile)
""")

doc.add_page_break()

# 4. Technology Stack
add_heading(doc, "4. TECHNOLOGY STACK", level=1)

table = create_table(doc, 1, 2, ["Category", "Technology"])
tech_data = [
    ("Framework", "React 18+"),
    ("Language", "TypeScript"),
    ("Build Tool", "Vite"),
    ("CSS Framework", "Tailwind CSS"),
    ("UI Component Library", "Shadcn/UI (Radix UI)"),
    ("Routing", "React Router DOM"),
    ("State Management", "React Context API + React Query (TanStack Query)"),
    ("Form Handling", "React Hook Form + Zod/Resolvers"),
    ("HTTP Client", "Fetch API"),
    ("Icons", "Lucide React"),
    ("Notifications", "Sonner (Toast notifications)"),
    ("Utilities", "clsx, class-variance-authority, date-fns"),
    ("Development", "ESLint, Tailwind CSS"),
    ("Containerization", "Docker"),
    ("Reverse Proxy", "Nginx"),
]
for category, tech in tech_data:
    row_cells = table.add_row().cells
    row_cells[0].text = category
    row_cells[1].text = tech

doc.add_page_break()

# 5. Project Structure
add_heading(doc, "5. PROJECT STRUCTURE", level=1)

add_heading(doc, "5.1 Detailed Directory Structure", level=2)
add_code_block(doc, """
src/
├── App.tsx                    # Root component with routing
├── main.tsx                   # Entry point
├── index.css                  # Global styles
├── App.css                    # App-specific styles
├── components/                # Reusable components
│   ├── AdminRoute.tsx         # Admin-only route guard
│   ├── AppSidebar.tsx         # Main navigation sidebar
│   ├── AssetManagement.tsx    # Asset CRUD interface
│   ├── AuthProvider.tsx       # Auth context provider
│   ├── CommentSystem.tsx      # Comment management
│   ├── Dashboard.tsx          # Main dashboard
│   ├── ErrorBoundary.tsx      # Global error handler
│   ├── Header.tsx             # Top navigation bar
│   ├── Layout.tsx             # Page layout wrapper
│   ├── PermissionRoute.tsx    # Permission-based route guard
│   ├── ProtectedRoute.tsx     # Auth-required route guard
│   ├── VulnerabilityManagement.tsx  # Vulnerability interface
│   ├── UserManagement.tsx     # User admin
│   ├── ui/                    # Shadcn/UI component library
│   │   ├── button.tsx
│   │   ├── dialog.tsx
│   │   ├── table.tsx
│   │   ├── card.tsx
│   │   ├── chart.tsx
│   │   └── ...40+ UI components
├── pages/                     # Route pages
│   ├── Index.tsx              # Dashboard page
│   ├── Assets.tsx             # Assets page
│   ├── Vulnerabilities.tsx    # Vulnerabilities page
│   ├── UserManagementPage.tsx # User management page
│   ├── Login.tsx              # Login page
│   ├── Register.tsx           # Registration page
│   ├── Settings.tsx           # Settings page
│   ├── Integrations.tsx       # Integrations page
│   └── ...15+ pages
├── services/                  # API and business logic
│   ├── integrationService.ts
│   └── microsoftSSOService.ts
├── hooks/                     # Custom React hooks
│   ├── use-toast.ts
│   └── use-mobile.tsx
├── contexts/                  # React context providers
├── config/                    # Configuration
│   └── environment.ts         # Environment variables
├── lib/                       # Utility libraries
│   └── utils.ts               # Helper functions
└── utils/                     # Utility functions
    └── tokenManager.ts        # JWT token management
""")

add_heading(doc, "5.2 Directory Purposes", level=2)

table = create_table(doc, 1, 2, ["Directory", "Purpose"])
dirs_data = [
    ("components/", "Reusable React components and layout components"),
    ("pages/", "Page components mapped to routes"),
    ("services/", "API calls and external service integration"),
    ("hooks/", "Custom React hooks for reusable logic"),
    ("contexts/", "React Context providers for global state"),
    ("config/", "Application configuration and constants"),
    ("lib/", "Utility functions and helper libraries"),
    ("utils/", "General utility functions (token management, etc.)"),
    ("ui/", "Shadcn/UI pre-built components"),
]
for dir_name, purpose in dirs_data:
    row_cells = table.add_row().cells
    row_cells[0].text = dir_name
    row_cells[1].text = purpose

doc.add_page_break()

# 6. Authentication & Authorization
add_heading(doc, "6. AUTHENTICATION & AUTHORIZATION", level=1)

add_heading(doc, "6.1 Authentication Overview", level=2)
add_paragraph(doc, """
Orbit-I implements a comprehensive authentication and authorization system:

• Authentication Methods: 
  - Credential-based login (username/password)
  - Microsoft SSO (Azure Active Directory)
  - JWT token-based session

• Token Storage: 
  - Secure SessionStorage (more secure than localStorage)
  - In-memory backup for quick access
  - Automatic expiry validation
""")

add_heading(doc, "6.2 AuthProvider Context", level=2)
add_paragraph(doc, """
The AuthProvider (src/components/AuthProvider.tsx) is the central authentication management component.

Key Responsibilities:
• Manage authentication state (isAuthenticated, token, userInfo)
• Persist user session across page reloads
• Validate token expiry
• Fetch user information from API
• Provide login/logout methods
• Check permissions and admin status

Context Properties:
""")

table = create_table(doc, 1, 2, ["Property", "Type & Description"])
auth_props = [
    ("isAuthenticated", "boolean - Current auth status"),
    ("token", "string | null - JWT token"),
    ("userInfo", "UserInfo | null - User details"),
    ("loading", "boolean - Initial load state"),
    ("login()", "Function to set token"),
    ("logout()", "Function to clear auth"),
    ("getToken()", "Function to get current token"),
    ("hasPermission()", "Function to check specific permission"),
    ("isAdmin()", "Function to check admin status"),
]
for prop, desc in auth_props:
    row_cells = table.add_row().cells
    row_cells[0].text = prop
    row_cells[1].text = desc

add_heading(doc, "6.3 Route Protection", level=2)
add_paragraph(doc, """
Three layers of route protection are implemented:

1. ProtectedRoute: Requires authentication
   - Redirects unauthenticated users to /login
   - Wraps all application routes

2. AdminRoute: Requires admin role
   - Only accessible to administrators
   - Wrapped inside ProtectedRoute

3. PermissionRoute: Fine-grained permissions
   - Requires specific permission string
   - Example: asset_management, integrations
   - Provides granular access control
""")

add_heading(doc, "6.4 Token Management (SecureTokenManager)", level=2)
add_paragraph(doc, """
SecureTokenManager (src/utils/tokenManager.ts) handles secure JWT token operations:

Methods:
• setToken(token, expiresIn): Store token with expiry time
• getToken(): Retrieve valid token (checks expiry)
• isTokenValid(): Check if current token is valid
• clearToken(): Remove token on logout
• refreshToken(refreshToken): Refresh expired token

Security Features:
• Uses SessionStorage (cleared when browser closes)
• Automatic expiry validation
• Prevents access with expired tokens
• Optional refresh token support
""")

doc.add_page_break()

# 7. Core Components
add_heading(doc, "7. CORE COMPONENTS", level=1)

add_heading(doc, "7.1 Layout Components", level=2)

add_heading(doc, "AppSidebar", level=3)
add_paragraph(doc, """
File: src/components/AppSidebar.tsx
Purpose: Main navigation sidebar showing menu items based on user role

Features:
• Role-based menu items (admin items only for admins)
• Collapsible on mobile
• Active route highlighting
• Integration with React Router
• Permission-aware items
""")

add_heading(doc, "Header", level=3)
add_paragraph(doc, """
File: src/components/Header.tsx
Purpose: Top navigation bar

Features:
• User profile dropdown
• Logout functionality
• App title/logo
• Responsive design
""")

add_heading(doc, "Layout", level=3)
add_paragraph(doc, """
File: src/components/Layout.tsx
Purpose: Wrapper component providing consistent page layout

Structure:
<Layout>
  └── SidebarProvider
      ├── AppSidebar (left navigation)
      └── Main content area
          ├── Header (top bar)
          └── Page content
""")

add_heading(doc, "7.2 Domain Components", level=2)

add_heading(doc, "VulnerabilityManagement", level=3)
add_paragraph(doc, """
File: src/components/VulnerabilityManagement.tsx (676 lines)
Purpose: Full vulnerability CRUD and management interface

Features:
• Paginated vulnerability table
• Search and filter (by name, CVE, severity)
• Severity-based color coding
• Export to CSV
• Exception creation workflow
• Comment system integration
• API integration with pagination
• Real-time severity counts

State Management:
• vulnerabilities: List of vulnerabilities
• filters: searchTerm, severityFilter, page, pageSize
• selectedVulnerability: Currently selected item
• exceptionFormData: Form state for exceptions
""")

add_heading(doc, "AssetManagement", level=3)
add_paragraph(doc, """
File: src/components/AssetManagement.tsx (1293 lines)
Purpose: Complete asset lifecycle management

Features:
• Create, Read, Update, Delete (CRUD) assets
• Bulk operations support
• Filter by status, owner, OS
• Rescan readiness tracking
• Owner and OS distribution statistics
• Paginated table with sorting
• Real-time validation

Asset Properties:
• id: Unique identifier
• hostname: Machine name
• host_ip: IP address
• last_scan_date: Last vulnerability scan
• status: online/offline/maintenance
• server_owner: Responsible team/person
• application_dependent: Supported applications
• os_name: Operating system
• ready_for_rescan: Rescan scheduled indicator
""")

add_heading(doc, "CommentSystem", level=3)
add_paragraph(doc, """
File: src/components/CommentSystem.tsx
Purpose: Add and display comments on vulnerabilities and hosts

Features:
• Add new comments with timestamps
• Display threaded comments
• User attribution
• Rich text support (optional)
• Real-time comment refresh
• Integration with vulnerability detail view
""")

add_heading(doc, "UserManagement", level=3)
add_paragraph(doc, """
File: src/components/UserManagement.tsx
Purpose: Admin interface for user management

Features:
• Create new users
• Edit user roles and permissions
• Delete users
• List all users with details
• Role assignment
• Permission management
""")

doc.add_page_break()

# 8. Pages & Routes
add_heading(doc, "8. PAGES & ROUTES", level=1)

add_heading(doc, "8.1 Route Map", level=2)

table = create_table(doc, 1, 3, ["Route", "Component", "Access"])
routes_data = [
    ("/login", "Login.tsx", "Public"),
    ("/register", "Register.tsx", "Public"),
    ("/auth/microsoft/callback", "MicrosoftCallback.tsx", "Public"),
    ("/", "Index.tsx", "Authenticated"),
    ("/vulnerabilities", "Vulnerabilities.tsx", "Authenticated"),
    ("/vulnerabilities/by-owner", "VulnerabilityByOwner.tsx", "Authenticated"),
    ("/vulnerabilities/by-hosts", "VulnerabilityByHosts.tsx", "Authenticated"),
    ("/vulnerabilities/closed", "ClosedVulnerabilitiesPage.tsx", "Authenticated"),
    ("/exceptions", "Exceptions.tsx", "Authenticated"),
    ("/hosts/:hostId", "HostDetails.tsx", "Authenticated"),
    ("/owners/:ownerName", "OwnerDetails.tsx", "Authenticated"),
    ("/assets", "Assets.tsx", "Permission: asset_management"),
    ("/scans", "ScanManagement.tsx", "Permission: scan_management"),
    ("/scan-history", "ScanHistory.tsx", "Authenticated"),
    ("/rescan-ready", "RescanReadyHosts.tsx", "Admin only"),
    ("/user-management", "UserManagementPage.tsx", "Admin only"),
    ("/integrations", "Integrations.tsx", "Permission: integrations"),
    ("/settings", "Settings.tsx", "Authenticated"),
    ("/*", "NotFound.tsx", "Public (404)"),
]
for route, component, access in routes_data:
    row_cells = table.add_row().cells
    row_cells[0].text = route
    row_cells[1].text = component
    row_cells[2].text = access

add_heading(doc, "8.2 Key Pages Description", level=2)

add_heading(doc, "Dashboard (Index.tsx)", level=3)
add_paragraph(doc, """
Main entry point for authenticated users.

Components:
• Header with user info
• Vulnerability summary cards
• Charts: Pie charts, trend charts, owner distribution
• Recent vulnerabilities table
• Quick action buttons
• Statistics overview
""")

add_heading(doc, "Vulnerabilities.tsx", level=3)
add_paragraph(doc, """
Primary vulnerability management interface.

Features:
• VulnerabilityManagement component integration
• Search and filtering
• Pagination
• Bulk operations
• Export functionality
""")

add_heading(doc, "Assets.tsx", level=3)
add_paragraph(doc, """
Infrastructure asset management page.

Features:
• AssetManagement component integration
• CRUD operations
• Filtering and searching
• Statistics display
""")

doc.add_page_break()

# 9. Services & API Integration
add_heading(doc, "9. SERVICES & API INTEGRATION", level=1)

add_heading(doc, "9.1 API Communication Pattern", level=2)
add_paragraph(doc, """
Orbit-I uses a custom fetch wrapper for API calls:

Base URL: Configured in environment.ts
• Development: http://orbiti.fareportal.com:7000
• Production: https://orbiti.fareportal.com:7000

All API calls include:
• Authorization header with JWT token
• Content-Type: application/json
• Error handling and user feedback
• Automatic token refresh on expiry
""")

add_heading(doc, "9.2 MicrosoftSSOService", level=2)
add_paragraph(doc, """
File: src/services/microsoftSSOService.ts

Purpose: Handle Microsoft Azure AD integration

Methods:
• getConfig(): Fetch SSO configuration from backend
• getLoginUrl(): Get Microsoft login URL
• handleCallback(code, state): Process OAuth callback
• generateState(): Generate CSRF protection token

Usage Example:
const ssoService = new MicrosoftSSOService();
const loginUrl = await ssoService.getLoginUrl();
window.location.href = loginUrl;
""")

add_heading(doc, "9.3 Integration Service", level=2)
add_paragraph(doc, """
File: src/services/integrationService.ts

Purpose: Manage third-party integrations

Likely Capabilities:
• List configured integrations
• Add new integrations
• Update integration settings
• Remove integrations
• Test integration connectivity
""")

add_heading(doc, "9.4 Custom Hooks for Data Fetching", level=2)
add_paragraph(doc, """
File: src/lib/utils.ts (useApiFetch hook)

Purpose: Centralized API communication with error handling

Features:
• Automatic token injection
• Error handling and user feedback
• Request/response logging
• Automatic retry logic
• Toast notifications on errors

Usage:
const apiFetch = useApiFetch();
const response = await apiFetch(`${config.apiUrl}/api/endpoint`);
const data = await response.json();
""")

doc.add_page_break()

# 10. State Management
add_heading(doc, "10. STATE MANAGEMENT", level=1)

add_heading(doc, "10.1 Context API (Global State)", level=2)
add_paragraph(doc, """
AuthProvider (src/components/AuthProvider.tsx)

Global State:
• isAuthenticated: boolean
• token: JWT token string
• userInfo: User object
• loading: Initial load state

Usage:
import { useContext } from 'react';
import { AuthContext } from './AuthProvider';

const { isAuthenticated, userInfo } = useContext(AuthContext);
""")

add_heading(doc, "10.2 React Query (Server State)", level=2)
add_paragraph(doc, """
TanStack Query is configured in App.tsx for server state management.

Benefits:
• Automatic caching of API responses
• Background refetching
• Stale-while-revalidate pattern
• Automatic garbage collection
• Optimistic updates support

Configuration (App.tsx):
const queryClient = new QueryClient();
<QueryClientProvider client={queryClient}>
  {/* App content */}
</QueryClientProvider>
""")

add_heading(doc, "10.3 Component Local State", level=2)
add_paragraph(doc, """
Individual components manage local state with useState:

Examples:
• VulnerabilityManagement: vulnerabilities list, filters, pagination
• AssetManagement: assets list, modals, selected item
• CommentSystem: form input, submission state

Pattern:
const [data, setData] = useState<Type>(initialValue);
const [isLoading, setIsLoading] = useState(false);
""")

doc.add_page_break()

# 11. UI Components Library
add_heading(doc, "11. UI COMPONENTS LIBRARY", level=1)

add_heading(doc, "11.1 Shadcn/UI Overview", level=2)
add_paragraph(doc, """
Orbit-I uses Shadcn/UI (built on Radix UI) for unstyled, accessible components.

Benefits:
• Accessibility built-in (WCAG compliant)
• Customizable with Tailwind CSS
• Small bundle size
• Composable architecture
• Type-safe with TypeScript

40+ Components Available:
• Form controls: button, input, select, textarea, checkbox, radio-group
• Layout: card, separator, tabs, accordion, collapsible
• Overlays: dialog, drawer, dropdown-menu, popover, tooltip
• Data display: table, avatar, badge, progress, skeleton
• Navigation: breadcrumb, pagination, tabs, navigation-menu
• Other: chart, calendar, carousel, command, context-menu
""")

add_heading(doc, "11.2 Custom UI Usage Patterns", level=2)

add_heading(doc, "Button Component", level=3)
add_code_block(doc, """
import { Button } from "@/components/ui/button";

<Button onClick={handleClick}>
  Click Me
</Button>

<Button variant="outline" size="sm">
  Outline Button
</Button>

<Button disabled>Disabled</Button>
""")

add_heading(doc, "Card Component", level=3)
add_code_block(doc, """
import { Card, CardContent, CardDescription, CardHeader, CardTitle } from "@/components/ui/card";

<Card>
  <CardHeader>
    <CardTitle>Title</CardTitle>
    <CardDescription>Description</CardDescription>
  </CardHeader>
  <CardContent>
    Content goes here
  </CardContent>
</Card>
""")

add_heading(doc, "Table Component", level=3)
add_code_block(doc, """
import { Table, TableBody, TableCell, TableHead, TableHeader, TableRow } from "@/components/ui/table";

<Table>
  <TableHeader>
    <TableRow>
      <TableHead>Header 1</TableHead>
      <TableHead>Header 2</TableHead>
    </TableRow>
  </TableHeader>
  <TableBody>
    <TableRow>
      <TableCell>Data 1</TableCell>
      <TableCell>Data 2</TableCell>
    </TableRow>
  </TableBody>
</Table>
""")

add_heading(doc, "Dialog Component", level=3)
add_code_block(doc, """
import { Dialog, DialogContent, DialogDescription, DialogHeader, DialogTitle, DialogTrigger } from "@/components/ui/dialog";

<Dialog>
  <DialogTrigger asChild>
    <Button>Open Dialog</Button>
  </DialogTrigger>
  <DialogContent>
    <DialogHeader>
      <DialogTitle>Dialog Title</DialogTitle>
      <DialogDescription>Dialog description</DialogDescription>
    </DialogHeader>
    {/* Content */}
  </DialogContent>
</Dialog>
""")

doc.add_page_break()

# 12. Styling & Theming
add_heading(doc, "12. STYLING & THEMING", level=1)

add_heading(doc, "12.1 Tailwind CSS Setup", level=2)
add_paragraph(doc, """
Tailwind CSS is used for utility-first styling.

Configuration: tailwind.config.ts

Key Features:
• Responsive breakpoints (mobile-first)
• Color system with semantic names
• Custom spacing scale
• Utility classes for rapid development
• CSS optimization for production

Breakpoints:
• sm: 640px
• md: 768px
• lg: 1024px
• xl: 1280px
• 2xl: 1536px

Usage Example:
<div className="p-4 md:p-6 bg-slate-100 rounded-lg shadow-md">
  Content with responsive padding
</div>
""")

add_heading(doc, "12.2 Color System", level=2)
add_paragraph(doc, """
Semantic colors used throughout the application:

• Severity Colors:
  - Critical: Red (rgb(220, 38, 38))
  - High: Orange (rgb(234, 88, 12))
  - Medium: Yellow (rgb(202, 138, 4))
  - Low: Green (rgb(34, 197, 94))

• Status Colors:
  - Active/Online: Green
  - Inactive/Offline: Gray
  - Pending: Yellow
  - Error: Red

• UI Colors:
  - Primary: Blue (for buttons, links)
  - Background: Light gray/white
  - Border: Light gray
  - Text: Dark gray/black
""")

add_heading(doc, "12.3 CSS Files", level=2)

table = create_table(doc, 1, 2, ["File", "Purpose"])
css_files = [
    ("index.css", "Global styles, Tailwind directives"),
    ("App.css", "App-specific styles"),
]
for file, purpose in css_files:
    row_cells = table.add_row().cells
    row_cells[0].text = file
    row_cells[1].text = purpose

doc.add_page_break()

# 13. Utilities & Hooks
add_heading(doc, "13. UTILITIES & HOOKS", level=1)

add_heading(doc, "13.1 Custom Hooks", level=2)

add_heading(doc, "use-toast Hook", level=3)
add_paragraph(doc, """
File: src/hooks/use-toast.ts

Purpose: Display toast notifications to users

Usage:
import { useToast } from '@/hooks/use-toast';

const { toast } = useToast();

toast({
  title: 'Success',
  description: 'Operation completed successfully',
  variant: 'default'
});

// For error
toast({
  title: 'Error',
  description: 'Something went wrong',
  variant: 'destructive'
});
""")

add_heading(doc, "use-mobile Hook", level=3)
add_paragraph(doc, """
File: src/hooks/use-mobile.tsx

Purpose: Detect mobile device and responsive breakpoints

Usage:
import { useIsMobile } from '@/hooks/use-mobile';

const isMobile = useIsMobile();

{isMobile && <MobileMenu />}
{!isMobile && <DesktopMenu />}
""")

add_heading(doc, "13.2 Utility Functions", level=2)

add_heading(doc, "utils.ts", level=3)
add_paragraph(doc, """
File: src/lib/utils.ts

Key Functions:
• cn(): Combine classnames with clsx
• useApiFetch(): Custom hook for API calls with error handling
• Helper functions for common operations

Example - cn() function:
import { cn } from '@/lib/utils';

const className = cn(
  'p-4 rounded-lg',
  isActive && 'bg-blue-500',
  isDisabled && 'opacity-50'
);
""")

add_heading(doc, "13.3 Security Utilities", level=2)

add_heading(doc, "tokenManager.ts", level=3)
add_paragraph(doc, """
File: src/utils/tokenManager.ts

Purpose: Secure JWT token management

Key Methods:
• setToken(token, expiresIn): Store token with expiry
• getToken(): Retrieve valid token
• isTokenValid(): Check token validity
• clearToken(): Remove token
• refreshToken(): Refresh expired token

Security Features:
• SessionStorage (cleared on browser close)
• Automatic expiry validation
• In-memory backup for performance
• Prevents use of expired tokens
""")

doc.add_page_break()

# 14. Development Workflow
add_heading(doc, "14. DEVELOPMENT WORKFLOW", level=1)

add_heading(doc, "14.1 Setup Instructions", level=2)
add_code_block(doc, """
# Prerequisites:
# - Node.js 16+
# - npm or yarn

# Clone repository
git clone <repository-url>
cd VMT-FRONTEND

# Install dependencies
npm install

# Set up environment variables
# Create .env file in root directory with:
VITE_API_URL=http://orbiti.fareportal.com:7000
VITE_ENABLE_HTTPS=false
VITE_TOKEN_EXPIRY=3600
VITE_ENABLE_DEBUG=true
""")

add_heading(doc, "14.2 Development Commands", level=2)

table = create_table(doc, 1, 2, ["Command", "Purpose"])
commands = [
    ("npm run dev", "Start development server (Vite)"),
    ("npm run build", "Build for production"),
    ("npm run build:dev", "Build in development mode"),
    ("npm run lint", "Run ESLint to check code quality"),
    ("npm run preview", "Preview production build locally"),
]
for cmd, purpose in commands:
    row_cells = table.add_row().cells
    row_cells[0].text = cmd
    row_cells[1].text = purpose

add_heading(doc, "14.3 Development Workflow", level=2)
add_paragraph(doc, """
1. Start dev server: npm run dev
2. Open http://localhost:8080 in browser
3. Make changes to components/pages
4. Changes auto-reload with HMR (Hot Module Replacement)
5. Check browser console for errors
6. Run linting: npm run lint
7. Commit changes: git commit -m "description"

Code Style:
• Use TypeScript for type safety
• Follow naming conventions (PascalCase for components, camelCase for functions)
• Add comments for complex logic
• Use meaningful variable names
• Keep components focused on single responsibility
""")

doc.add_page_break()

# 15. Build & Deployment
add_heading(doc, "15. BUILD & DEPLOYMENT", level=1)

add_heading(doc, "15.1 Build Process", level=2)
add_code_block(doc, """
# Production build
npm run build

# Output:
# dist/
# ├── index.html
# ├── assets/
# │   ├── index-xxx.js     (JavaScript bundle)
# │   ├── index-xxx.css    (CSS bundle)
# │   └── ...
# └── ...

# Build outputs minified, optimized code for production
# CSS is purged of unused styles
# JavaScript is tree-shaken for smaller bundle size
""")

add_heading(doc, "15.2 Docker Deployment", level=2)
add_paragraph(doc, """
Docker configuration is provided for containerized deployment.

Files:
• Dockerfile: Production image definition
• docker-compose.yml: Multi-container orchestration
• nginx.conf: Development reverse proxy
• nginx-ssl.conf: Production reverse proxy with SSL

Docker Build:
docker build -t orbit-i-frontend:latest .

Docker Run:
docker run -p 8080:80 orbit-i-frontend:latest

Docker Compose:
docker-compose up -d
""")

add_heading(doc, "15.3 SSL/HTTPS Configuration", level=2)
add_paragraph(doc, """
Directories:
• cert/: Development SSL certificates
• certs/: Production SSL certificates

Nginx Configuration:
• nginx.conf: HTTP reverse proxy for development
• nginx-ssl.conf: HTTPS with SSL certificates for production

To use SSL:
1. Place SSL certificate and key in certs/ directory
2. Use nginx-ssl.conf in Dockerfile
3. Rebuild and deploy
""")

add_heading(doc, "15.4 Environment-Specific Configuration", level=2)
add_paragraph(doc, """
Environment variables (configured via .env file):

VITE_API_URL
• Dev: http://orbiti.fareportal.com:7000
• Prod: https://orbiti.fareportal.com:7000

VITE_ENABLE_HTTPS
• Dev: false
• Prod: true

VITE_TOKEN_EXPIRY
• Default: 3600 seconds (1 hour)
• Configurable per environment

VITE_ENABLE_DEBUG
• Dev: true (enable debug logging)
• Prod: false

VITE_ENABLE_ANALYTICS
• Enable/disable analytics tracking
""")

doc.add_page_break()

# 16. Security Considerations
add_heading(doc, "16. SECURITY CONSIDERATIONS", level=1)

add_heading(doc, "16.1 Authentication Security", level=2)
add_paragraph(doc, """
Implemented Security Measures:
• JWT tokens stored in SessionStorage (not localStorage)
• Automatic token expiry validation
• Tokens cleared on logout
• Secure refresh token mechanism (when available)
• HTTPS required in production
• CORS configuration on backend
""")

add_heading(doc, "16.2 Authorization Security", level=2)
add_paragraph(doc, """
Access Control:
• Role-based access (admin, team, user)
• Permission-based access (fine-grained)
• Route-level protection (ProtectedRoute, AdminRoute, PermissionRoute)
• Server-side validation required for all operations

Best Practices:
• Never trust client-side authorization alone
• Always validate permissions on backend
• Use role names from authenticated user
• Check permissions before API calls (UX)
• Implement backend permission checks (security)
""")

add_heading(doc, "16.3 Data Security", level=2)
add_paragraph(doc, """
Protected Data:
• Authentication tokens (SessionStorage)
• User credentials (transmitted over HTTPS only)
• Sensitive user information (never logged)
• API keys (never exposed in frontend code)

Never:
• Store secrets in environment variables that get bundled
• Log sensitive data to console in production
• Commit .env files with secrets
• Expose API keys in client-side code
""")

add_heading(doc, "16.4 Input Validation", level=2)
add_paragraph(doc, """
Frontend Validation:
• React Hook Form validates form inputs
• TypeScript provides type safety
• Input sanitization before display

Backend Validation Required:
• Never trust frontend validation alone
• Backend must validate all inputs
• SQL injection prevention (parameterized queries)
• XSS protection (output encoding)
• CSRF protection (token validation)
""")

doc.add_page_break()

# 17. Performance Optimization
add_heading(doc, "17. PERFORMANCE OPTIMIZATION", level=1)

add_heading(doc, "17.1 Frontend Optimization Techniques", level=2)

add_heading(doc, "Code Splitting", level=3)
add_paragraph(doc, """
Vite automatically code-splits the application:
• Main bundle includes core application
• Lazy-loaded route components
• Component-level splitting for large components

Benefit: Faster initial page load
""")

add_heading(doc, "React Query Caching", level=3)
add_paragraph(doc, """
Automatic caching reduces API calls:
• Cached responses reused for same queries
• Stale-while-revalidate for background updates
• Automatic garbage collection of unused data

Configuration:
const queryClient = new QueryClient({
  defaultOptions: {
    queries: {
      staleTime: 1000 * 60 * 5, // 5 minutes
      cacheTime: 1000 * 60 * 10, // 10 minutes
    },
  },
});
""")

add_heading(doc, "Component Memoization", level=3)
add_paragraph(doc, """
Prevent unnecessary re-renders:
• Use React.memo for presentational components
• useMemo for expensive computations
• useCallback for function props

Example:
const MemoizedComponent = React.memo(MyComponent);
""")

add_heading(doc, "17.2 Bundle Size Optimization", level=2)
add_paragraph(doc, """
Vite provides:
• Tree-shaking: Remove unused code
• Minification: Reduce code size
• CSS purging: Remove unused styles
• Asset optimization: Compress images

Production Build Size:
• Main JS: ~200-300 KB (gzipped ~50-70 KB)
• CSS: ~50-100 KB (gzipped ~10-20 KB)
• Total: Depends on components used
""")

add_heading(doc, "17.3 Network Performance", level=2)
add_paragraph(doc, """
Optimization Strategies:
• Use HTTPS/HTTP2 for compression
• Enable Gzip compression on server
• Use CDN for static assets
• Implement lazy loading for images
• Minimize API calls with React Query
• Implement request batching if possible
""")

doc.add_page_break()

# 18. Testing Strategy
add_heading(doc, "18. TESTING STRATEGY", level=1)

add_heading(doc, "18.1 Testing Recommendations", level=2)
add_paragraph(doc, """
Recommended Testing Approach:

Unit Tests:
• Test utility functions
• Test custom hooks
• Test service classes
• Tools: Vitest, Jest

Component Tests:
• Test component rendering
• Test user interactions
• Test props and state
• Tools: Vitest + React Testing Library

Integration Tests:
• Test component interactions
• Test routing
• Test API integration
• Tools: Cypress, Playwright

E2E Tests:
• Test complete user workflows
• Test authentication flow
• Test real API integration
• Tools: Cypress, Playwright
""")

add_heading(doc, "18.2 Testing Best Practices", level=2)
add_paragraph(doc, """
• Test user behavior, not implementation
• Keep tests simple and focused
• Mock external dependencies (API calls)
• Test error cases and edge cases
• Maintain >80% code coverage for critical paths
• Run tests before committing code
• Use CI/CD for automated testing
""")

doc.add_page_break()

# 19. Troubleshooting & FAQ
add_heading(doc, "19. TROUBLESHOOTING & FAQ", level=1)

add_heading(doc, "19.1 Common Issues & Solutions", level=2)

add_heading(doc, "Issue: Blank Page on Load", level=3)
add_paragraph(doc, """
Causes:
• API URL not configured
• Authentication token invalid
• JavaScript error in console

Solutions:
1. Check .env file for VITE_API_URL
2. Clear browser cache and cookies
3. Check browser console for errors (F12)
4. Verify backend API is running
5. Check network tab for failed requests
""")

add_heading(doc, "Issue: Authentication Fails", level=3)
add_paragraph(doc, """
Causes:
• Wrong API URL
• Backend authentication service down
• CORS error
• Token storage issue

Solutions:
1. Verify API URL in environment.ts
2. Check backend API logs
3. Check browser console for CORS errors
4. Clear SessionStorage: sessionStorage.clear()
5. Try incognito/private browsing mode
6. Verify Microsoft SSO configuration (if using)
""")

add_heading(doc, "Issue: Components Not Rendering", level=3)
add_paragraph(doc, """
Causes:
• Import path errors
• TypeScript compilation errors
• Missing dependencies

Solutions:
1. Check console for TypeScript errors
2. Verify import paths use @/ alias
3. Run npm install to update dependencies
4. Restart dev server: npm run dev
5. Check component export statements
""")

add_heading(doc, "Issue: Styling Not Applied", level=3)
add_paragraph(doc, """
Causes:
• Tailwind CSS not configured
• Conflicting styles
• CSS not imported

Solutions:
1. Verify tailwind.config.ts exists
2. Check index.css includes @tailwind directives
3. Clear browser cache
4. Restart dev server
5. Check browser DevTools (Inspect element)
6. Verify class names are valid Tailwind classes
""")

add_heading(doc, "Issue: API Calls Failing", level=3)
add_paragraph(doc, """
Causes:
• Backend API down
• CORS not configured
• Invalid token
• Network error

Solutions:
1. Check backend API is running
2. Verify VITE_API_URL is correct
3. Check network tab for request details
4. Check response status and body
5. Verify token is valid
6. Check backend logs for errors
7. Test API endpoint with curl/Postman
""")

add_heading(doc, "19.2 Frequently Asked Questions", level=2)

add_heading(doc, "Q: How do I add a new page?", level=3)
add_paragraph(doc, """
A: 1. Create component in src/pages/MyPage.tsx
   2. Add import in src/App.tsx
   3. Add route: <Route path="/my-page" element={<MyPage />} />
   4. Add navigation link in AppSidebar if needed
""")

add_heading(doc, "Q: How do I add a new component?", level=3)
add_paragraph(doc, """
A: 1. Create component in src/components/MyComponent.tsx
   2. Export component: export function MyComponent() { ... }
   3. Import in parent: import { MyComponent } from '@/components/MyComponent'
   4. Use: <MyComponent prop={value} />
""")

add_heading(doc, "Q: How do I call an API?", level=3)
add_paragraph(doc, """
A: 1. Get apiFetch hook: const apiFetch = useApiFetch();
   2. Make request: const response = await apiFetch(`${config.apiUrl}/api/endpoint`);
   3. Parse: const data = await response.json();
   4. Handle errors in try/catch block
""")

add_heading(doc, "Q: How do I check if user is admin?", level=3)
add_paragraph(doc, """
A: import { useAuth } from '@/components/AuthProvider';
   const { isAdmin } = useAuth();
   if (isAdmin()) { /* admin only code */ }
""")

add_heading(doc, "Q: How do I show a toast notification?", level=3)
add_paragraph(doc, """
A: import { useToast } from '@/hooks/use-toast';
   const { toast } = useToast();
   toast({ title: 'Success', description: 'Done!' });
""")

add_heading(doc, "Q: How do I protect a route?", level=3)
add_paragraph(doc, """
A: For authentication: <ProtectedRoute><Page /></ProtectedRoute>
   For admin: <AdminRoute><AdminPage /></AdminRoute>
   For permissions: <PermissionRoute requiredPermission="name"><Page /></PermissionRoute>
""")

doc.add_page_break()

# Additional Resources
add_heading(doc, "ADDITIONAL RESOURCES", level=1)

add_heading(doc, "Documentation & Links", level=2)
add_paragraph(doc, """
• React: https://react.dev
• TypeScript: https://www.typescriptlang.org
• Vite: https://vitejs.dev
• Tailwind CSS: https://tailwindcss.com
• Shadcn/UI: https://ui.shadcn.com
• React Router: https://reactrouter.com
• React Query: https://tanstack.com/query
• React Hook Form: https://react-hook-form.com
• Lucide Icons: https://lucide.dev
""")

add_heading(doc, "Development Tools", level=2)
add_paragraph(doc, """
• VS Code: Code editor
• VS Code Extensions:
  - ESLint
  - Prettier
  - Tailwind CSS IntelliSense
  - Thunder Client (API testing)
  - Dev Tools (browser extension)
""")

# Save document
output_path = "Orbit-I-Frontend-Documentation.docx"
doc.save(output_path)
print(f"✓ Documentation generated successfully: {output_path}")
